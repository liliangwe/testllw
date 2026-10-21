import csv
import io
from datetime import date

from openpyxl import Workbook, load_workbook

HEADERS = ["employee_no", "name", "department", "position", "email", "phone", "hire_date", "status"]
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"


def _normalize(rec: dict) -> dict:
    for k in ("employee_no", "name", "department", "position", "email", "phone", "status"):
        v = rec.get(k)
        rec[k] = str(v).strip() if v is not None else ""
    hd = rec.get("hire_date")
    if isinstance(hd, date):
        rec["hire_date"] = hd
    elif isinstance(hd, str) and hd.strip():
        rec["hire_date"] = date.fromisoformat(hd.strip())
    else:
        rec["hire_date"] = None
    rec["status"] = rec.get("status") or "active"
    return rec


def _rows_to_records(rows: list) -> list[dict]:
    if not rows:
        return []
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if header != HEADERS:
        raise ValueError(f"表头不匹配，应为 {HEADERS}")
    out = []
    for row in rows[1:]:
        if row is None or all(c in (None, "") for c in row):
            continue
        out.append(_normalize(dict(zip(HEADERS, row))))
    return out


def build_template(fmt: str = "xlsx") -> tuple[bytes, str]:
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(HEADERS)
        writer.writerow(["E001", "张三", "研发部", "工程师", "zhangsan@example.com", "13800000000", "2024-01-01", "active"])
        return buf.getvalue().encode("utf-8-sig"), CSV_MIME
    wb = Workbook()
    ws = wb.active
    ws.title = "员工导入"
    ws.append(HEADERS)
    ws.append(["E001", "张三", "研发部", "工程师", "zhangsan@example.com", "13800000000", "2024-01-01", "active"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), XLSX_MIME


def parse_workbook(data: bytes) -> list[dict]:
    wb = load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    return _rows_to_records(list(ws.iter_rows(values_only=True)))


def parse_csv(data: bytes) -> list[dict]:
    text = data.decode("utf-8-sig")
    return _rows_to_records(list(csv.reader(io.StringIO(text))))


def parse_file(data: bytes, fmt: str) -> list[dict]:
    return parse_csv(data) if fmt == "csv" else parse_workbook(data)